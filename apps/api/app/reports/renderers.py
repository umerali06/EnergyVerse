"""Deterministic finalized-report renderers for PDF, DOCX, and XLSX."""

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.entities import GeneratedReport

NAVY = "002865"
ORANGE = "FB4402"
LIGHT = "F2F4F7"


def _text(value: object) -> str:
    return str(value).replace("\u2011", "-").replace("\u2013", "-").replace("\u2014", "-")


def _snapshot_rows(value: Any, prefix: str = "") -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    if isinstance(value, dict):
        for key, child in value.items():
            label = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_snapshot_rows(child, label))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            rows.extend(_snapshot_rows(child, f"{prefix}[{index}]"))
    else:
        rows.append((prefix, "" if value is None else _text(value)))
    return rows


def _metadata(report: GeneratedReport) -> list[tuple[str, str]]:
    return [
        ("Report ID", report.id),
        ("Type", report.report_type.replace("_", " ").title()),
        ("Status", report.status.title()),
        ("Generated", report.created_at.isoformat()),
        ("Finalized", report.finalized_at.isoformat() if report.finalized_at else ""),
        ("AI model", report.ai_model),
    ]


def render_pdf(report: GeneratedReport) -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=LETTER,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title=_text(report.title),
        author="Flacron EnergyVerse",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "FevTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=27,
        textColor=colors.HexColor(f"#{NAVY}"),
        alignment=TA_CENTER,
        spaceAfter=16,
    )
    heading = ParagraphStyle(
        "FevHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor(f"#{NAVY}"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body = ParagraphStyle(
        "FevBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        spaceAfter=6,
    )
    story: list[Any] = [
        Paragraph(
            "FLACRON ENERGYVERSE",
            ParagraphStyle(
                "Kicker",
                parent=body,
                textColor=colors.HexColor(f"#{ORANGE}"),
                alignment=TA_CENTER,
                fontSize=9,
                leading=11,
            ),
        ),
        Spacer(1, 6),
        Paragraph(_text(report.title), title),
    ]
    meta = Table(_metadata(report), colWidths=[1.35 * inch, 4.65 * inch])
    meta.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{LIGHT}")),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor(f"#{NAVY}")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D6DCE5")),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.extend(
        [
            meta,
            Paragraph("Executive Summary", heading),
            Paragraph(_text(report.narrative.summary), body),
        ]
    )
    for label, items in (
        ("Findings", report.narrative.findings),
        ("Recommendations", report.narrative.recommendations),
    ):
        story.append(Paragraph(label, heading))
        story.extend(Paragraph(f"- {_text(item)}", body) for item in items)
    if report.narrative.risk_score is not None:
        story.extend(
            [
                Paragraph("Risk Score", heading),
                Paragraph(f"{report.narrative.risk_score:.1f} / 100", body),
            ]
        )
    story.extend([PageBreak(), Paragraph("Authoritative Source Snapshot", heading)])
    rows = [["Field", "Value"], *[list(row) for row in _snapshot_rows(report.source_snapshot)]]
    table = Table(rows, colWidths=[2.2 * inch, 3.8 * inch], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DCE5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F8FAFC")],
                ),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return output.getvalue()


def _set_font(run: Any, size: float, color: str = "111827", bold: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def render_docx(report: GeneratedReport) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = (
        Inches(1)
    )
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(6), 1.1
    for name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6)):
        style = document.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = (
            "Calibri",
            Pt(size),
            RGBColor.from_string(NAVY),
        )
        style.paragraph_format.space_before, style.paragraph_format.space_after = (
            Pt(before),
            Pt(after),
        )
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(header.add_run("Flacron EnergyVerse | Finalized Report"), 8.5, "6B7280")
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(kicker.add_run("FLACRON ENERGYVERSE"), 9, ORANGE, True)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(16)
    _set_font(title.add_run(_text(report.title)), 22, NAVY, True)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for label, value in _metadata(report):
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Inches(1.35), Inches(5.15)
        _set_font(cells[0].paragraphs[0].add_run(label), 9, NAVY, True)
        _set_font(cells[1].paragraphs[0].add_run(_text(value)), 9)
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(_text(report.narrative.summary))
    for label, items in (
        ("Findings", report.narrative.findings),
        ("Recommendations", report.narrative.recommendations),
    ):
        document.add_heading(label, level=1)
        for item in items:
            document.add_paragraph(_text(item), style="List Bullet")
    if report.narrative.risk_score is not None:
        document.add_heading("Risk Score", level=1)
        document.add_paragraph(f"{report.narrative.risk_score:.1f} / 100")
    document.add_page_break()  # type: ignore[no-untyped-call]
    document.add_heading("Authoritative Source Snapshot", level=1)
    source = document.add_table(rows=1, cols=2)
    source.style = "Table Grid"
    source.autofit = False
    source.rows[0].cells[0].text, source.rows[0].cells[1].text = "Field", "Value"
    source.rows[0].cells[0].width, source.rows[0].cells[1].width = Inches(2.2), Inches(4.3)
    for cell in source.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            _set_font(run, 9, "FFFFFF", True)
        cell._tc.get_or_add_tcPr().append(_cell_fill(NAVY))
    for field, value in _snapshot_rows(report.source_snapshot):
        cells = source.add_row().cells
        cells[0].width, cells[1].width = Inches(2.2), Inches(4.3)
        _set_font(cells[0].paragraphs[0].add_run(field), 8)
        _set_font(cells[1].paragraphs[0].add_run(value), 8)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _cell_fill(color: str) -> Any:
    from docx.oxml import OxmlElement

    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), color)
    return shade


def render_xlsx(report: GeneratedReport) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert isinstance(summary, Worksheet)
    summary.title = "Report Summary"
    summary.sheet_view.showGridLines = False
    summary.freeze_panes = "A4"
    summary.page_setup.orientation = "portrait"
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary_page_setup = summary.sheet_properties.pageSetUpPr
    assert summary_page_setup is not None
    summary_page_setup.fitToPage = True
    summary_page_setup.autoPageBreaks = False
    summary.print_options.horizontalCentered = True
    summary.merge_cells("A1:D1")
    summary["A1"] = _text(report.title)
    summary["A1"].font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.row_dimensions[1].height = 32
    row = 3
    for label, value in _metadata(report):
        summary.cell(row, 1, label).font = Font(bold=True, color=NAVY)
        summary.cell(row, 2, _text(value))
        row += 1
    row += 1
    for heading, content in (
        ("Executive Summary", [report.narrative.summary]),
        ("Findings", report.narrative.findings),
        ("Recommendations", report.narrative.recommendations),
    ):
        summary.cell(row, 1, heading).font = Font(size=12, bold=True, color=NAVY)
        summary.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT)
        row += 1
        for item in content:
            summary.cell(row, 1, _text(item))
            summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            summary.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 1
    if report.narrative.risk_score is not None:
        summary.cell(row, 1, "Risk Score").font = Font(bold=True, color=NAVY)
        summary.cell(row, 2, report.narrative.risk_score).number_format = '0.0 " / 100"'
    for column, width in {"A": 24, "B": 28, "C": 18, "D": 18}.items():
        summary.column_dimensions[column].width = width

    source = workbook.create_sheet("Source Snapshot")
    source.sheet_view.showGridLines = False
    source.freeze_panes = "A2"
    source.page_setup.orientation = "landscape"
    source.page_setup.fitToWidth = 1
    source.page_setup.fitToHeight = 0
    source_page_setup = source.sheet_properties.pageSetUpPr
    assert source_page_setup is not None
    source_page_setup.fitToPage = True
    source_page_setup.autoPageBreaks = False
    source.print_title_rows = "1:1"
    source.append(["Field", "Value"])
    for field, value in _snapshot_rows(report.source_snapshot):
        source.append([field, value])
    for cell in source[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    thin = Side(style="thin", color="D6DCE5")
    for cells in source.iter_rows(min_row=2):
        for cell in cells:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    source.column_dimensions[get_column_letter(1)].width = 42
    source.column_dimensions[get_column_letter(2)].width = 80
    source.auto_filter.ref = f"A1:B{source.max_row}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


RENDERERS = {"pdf": render_pdf, "docx": render_docx, "xlsx": render_xlsx}
CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
