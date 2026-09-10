import asyncio
import uuid
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.audit.service import AuditService
from app.auth.dependencies import get_current_user
from app.db.repositories.assets import AssetRepository
from app.db.repositories.audit_logs import AuditLogRepository
from app.db.repositories.companies import CompanyRepository
from app.db.repositories.generated_reports import GeneratedReportRepository
from app.db.repositories.inspections import InspectionRepository
from app.db.repositories.safety_reports import SafetyReportRepository
from app.db.repositories.work_orders import WorkOrderRepository
from app.main import app
from app.models.entities import CurrentUser, ReportNarrative
from app.rbac.dependencies import get_access_denial_audit
from app.reports.service import GeneratedReportService, get_generated_report_service
from scripts.seed import ACME_COMPANY_ID, DEMO_INSPECTIONS, run_seed
from tests.fakes.firestore import FakeAsyncClient

BETA_COMPANY_ID = "beta-utilities"
COMPLETED_INSPECTION_ID = DEMO_INSPECTIONS[0].id


class FakeNarrativeClient:
    model_name = "test-report-model"

    async def generate(self, report_type: str, source_snapshot: dict[str, Any]) -> ReportNarrative:
        source = source_snapshot.get("source", {})
        return ReportNarrative(
            summary=f"Advisory {report_type} summary for {source.get('id', 'tenant')}",
            findings=["Source data was captured."],
            recommendations=["Human reviewer must verify this draft."],
            risk_score=25,
        )


class FakeReportStorage:
    def __init__(self) -> None:
        self.objects: dict[str, tuple[bytes, str]] = {}

    def upload(
        self, company_id: str, report_id: str, filename: str, data: bytes, content_type: str
    ) -> str:
        path = f"companies/{company_id}/reports/{report_id}/exports/{filename}"
        self.objects[path] = (data, content_type)
        return path

    def signed_url_for(self, path: str) -> str:
        return f"https://signed.invalid/{path}"


@pytest.fixture()
def wiring() -> dict[str, Any]:
    client = FakeAsyncClient()
    asyncio.run(run_seed(client))
    audit = AuditService(AuditLogRepository(client))
    storage = FakeReportStorage()
    service = GeneratedReportService(
        reports=GeneratedReportRepository(client, audit),
        companies=CompanyRepository(client, audit),
        assets=AssetRepository(client, audit),
        inspections=InspectionRepository(client, audit),
        work_orders=WorkOrderRepository(client, audit),
        safety_reports=SafetyReportRepository(client, audit),
        narrative_client=FakeNarrativeClient(),
        storage=storage,
    )
    app.dependency_overrides[get_generated_report_service] = lambda: service
    app.dependency_overrides[get_access_denial_audit] = lambda: audit
    yield {"client": client, "storage": storage}
    app.dependency_overrides.pop(get_generated_report_service, None)
    app.dependency_overrides.pop(get_access_denial_audit, None)


def _identity(
    *,
    uid: str = "report-author",
    company_id: str = ACME_COMPANY_ID,
    permissions: frozenset[str] = frozenset(
        {
            "reports.read",
            "reports.generate",
            "inspections.read",
            "assets.read",
            "work_orders.read",
            "safety.read",
        }
    ),
) -> CurrentUser:
    return CurrentUser(
        uid=uid,
        email=f"{uid}@example.invalid",
        email_verified=True,
        company_id=company_id,
        company_name="Test Co",
        role_key="custom",
        permissions=permissions,
    )


def _request(identity: CurrentUser, method: str, path: str, **kwargs: Any) -> Any:
    app.dependency_overrides[get_current_user] = lambda: identity
    try:
        with TestClient(app) as client:
            return client.request(method, path, **kwargs)
    finally:
        app.dependency_overrides.pop(get_current_user, None)


def _generate(identity: CurrentUser, **overrides: Any) -> Any:
    payload: dict[str, Any] = {
        "id": str(uuid.uuid4()),
        "report_type": "executive_summary",
    }
    payload.update(overrides)
    return _request(identity, "POST", "/api/v1/reports/generate", json=payload)


def test_generate_inspection_report_snapshots_source_and_audits(
    wiring: dict[str, Any],
) -> None:
    response = _generate(_identity(), report_type="inspection", source_id=COMPLETED_INSPECTION_ID)
    assert response.status_code == 200, response.json()
    body = response.json()
    assert body["status"] == "draft"
    assert body["source_snapshot"]["source"]["id"] == COMPLETED_INSPECTION_ID
    assert body["source_revision"] == body["source_snapshot"]["source"]["revision"]
    assert body["ai_model"] == "test-report-model"
    assert body["narrative"]["summary"].startswith("Advisory inspection")

    audits = wiring["client"]._store["audit_logs"].values()
    assert any(
        event["action"] == "generated_report.created" and event["target_id"] == body["id"]
        for event in audits
    )


def test_generate_requires_source_domain_permission(wiring: dict[str, Any]) -> None:
    identity = _identity(permissions=frozenset({"reports.read", "reports.generate"}))
    response = _generate(identity, report_type="inspection", source_id=COMPLETED_INSPECTION_ID)
    assert response.status_code == 403
    assert response.json()["error"] == "source_permission_required"


def test_report_reads_are_tenant_isolated(wiring: dict[str, Any]) -> None:
    created = _generate(_identity()).json()
    response = _request(
        _identity(company_id=BETA_COMPANY_ID), "GET", f"/api/v1/reports/{created['id']}"
    )
    assert response.status_code == 404
    assert response.json()["error"] == "generated_report_not_found"


def test_draft_edit_then_finalization_is_revision_safe_and_immutable(
    wiring: dict[str, Any],
) -> None:
    identity = _identity()
    created = _generate(identity).json()
    updated = _request(
        identity,
        "PATCH",
        f"/api/v1/reports/{created['id']}",
        json={
            "summary": "Human-reviewed executive summary.",
            "expected_revision": created["revision"],
        },
    )
    assert updated.status_code == 200
    assert updated.json()["narrative"]["summary"] == "Human-reviewed executive summary."

    stale = _request(
        identity,
        "POST",
        f"/api/v1/reports/{created['id']}/finalize",
        json={"expected_revision": created["revision"], "finalization_attestation": True},
    )
    assert stale.status_code == 409
    assert stale.json()["error"] == "revision_conflict"

    finalized = _request(
        identity,
        "POST",
        f"/api/v1/reports/{created['id']}/finalize",
        json={
            "expected_revision": updated.json()["revision"],
            "finalization_attestation": True,
        },
    )
    assert finalized.status_code == 200
    assert finalized.json()["status"] == "finalized"
    assert finalized.json()["finalized_by"] == identity.uid

    edit = _request(
        identity,
        "PATCH",
        f"/api/v1/reports/{created['id']}",
        json={
            "summary": "Attempted rewrite",
            "expected_revision": finalized.json()["revision"],
        },
    )
    assert edit.status_code == 409
    assert edit.json()["error"] == "finalized_report_immutable"
    delete = _request(identity, "DELETE", f"/api/v1/reports/{created['id']}")
    assert delete.status_code == 409


def test_inspection_with_unreviewed_ai_cannot_finalize(wiring: dict[str, Any]) -> None:
    now = datetime.now(UTC)
    inspection = wiring["client"]._store["inspections"][COMPLETED_INSPECTION_ID]
    inspection["ai_analysis"] = [
        {
            "id": "analysis-1",
            "media_local_id": "photo-1",
            "model": "test-vision",
            "summary": "Possible corrosion",
            "annotation_ids": [],
            "reviewed": False,
            "created_by": "inspector",
            "created_at": now,
        }
    ]
    created = _generate(
        _identity(), report_type="inspection", source_id=COMPLETED_INSPECTION_ID
    ).json()
    response = _request(
        _identity(),
        "POST",
        f"/api/v1/reports/{created['id']}/finalize",
        json={"expected_revision": created["revision"], "finalization_attestation": True},
    )
    assert response.status_code == 409
    assert response.json()["error"] == "unreviewed_ai_findings"


def test_regenerate_replaces_draft_with_current_source_snapshot(
    wiring: dict[str, Any],
) -> None:
    identity = _identity()
    created = _generate(
        identity, report_type="inspection", source_id=COMPLETED_INSPECTION_ID
    ).json()
    inspection = wiring["client"]._store["inspections"][COMPLETED_INSPECTION_ID]
    inspection["title"] = "Updated inspection title"
    inspection["revision"] += 1

    regenerated = _request(
        identity,
        "POST",
        f"/api/v1/reports/{created['id']}/regenerate",
        json={"expected_revision": created["revision"]},
    )
    assert regenerated.status_code == 200, regenerated.json()
    assert regenerated.json()["source_snapshot"]["source"]["title"] == ("Updated inspection title")
    assert regenerated.json()["source_revision"] == inspection["revision"]
    assert regenerated.json()["revision"] == created["revision"] + 1

    audits = wiring["client"]._store["audit_logs"].values()
    assert any(
        event["action"] == "generated_report.regenerated" and event["target_id"] == created["id"]
        for event in audits
    )


def test_report_list_filters_and_paginates(wiring: dict[str, Any]) -> None:
    identity = _identity()
    for _ in range(3):
        assert _generate(identity).status_code == 200
    first = _request(
        identity,
        "GET",
        "/api/v1/reports",
        params={"report_type": "executive_summary", "limit": 2},
    )
    assert first.status_code == 200
    assert len(first.json()["items"]) == 2
    assert first.json()["next_cursor"] is not None
    second = _request(
        identity,
        "GET",
        "/api/v1/reports",
        params={"report_type": "executive_summary", "cursor": first.json()["next_cursor"]},
    )
    assert second.status_code == 200
    assert len(second.json()["items"]) == 1


def test_read_and_generate_permissions_are_independent(wiring: dict[str, Any]) -> None:
    read_only = _identity(permissions=frozenset({"reports.read"}))
    assert _request(read_only, "GET", "/api/v1/reports").status_code == 200
    assert _generate(read_only).status_code == 403

    generate_only = _identity(permissions=frozenset({"reports.generate", "reports.read"}))
    assert _generate(generate_only).status_code == 200


def test_only_finalized_reports_export_to_private_valid_artifacts(
    wiring: dict[str, Any],
) -> None:
    identity = _identity()
    created = _generate(identity).json()
    draft = _request(
        identity, "POST", f"/api/v1/reports/{created['id']}/export", params={"format": "pdf"}
    )
    assert draft.status_code == 409
    assert draft.json()["error"] == "report_not_finalized"
    finalized = _request(
        identity,
        "POST",
        f"/api/v1/reports/{created['id']}/finalize",
        json={"expected_revision": created["revision"], "finalization_attestation": True},
    ).json()

    for export_format in ("pdf", "docx", "xlsx"):
        response = _request(
            identity,
            "POST",
            f"/api/v1/reports/{created['id']}/export",
            params={"format": export_format},
        )
        assert response.status_code == 200, response.json()
        body = response.json()
        assert body["url"].startswith("https://signed.invalid/companies/")
        assert body["size"] > 500
        path = body["url"].removeprefix("https://signed.invalid/")
        data = wiring["storage"].objects[path][0]
        if export_format == "pdf":
            assert "Executive Summary" in "".join(
                page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages
            )
        elif export_format == "docx":
            document = Document(BytesIO(data))
            assert any("Executive Summary" in p.text for p in document.paragraphs)
        else:
            workbook = load_workbook(BytesIO(data), data_only=False)
            assert workbook.sheetnames == ["Report Summary", "Source Snapshot"]
            assert workbook["Report Summary"]["A1"].value == finalized["title"]

    stored = _request(identity, "GET", f"/api/v1/reports/{created['id']}").json()
    assert stored["status"] == "finalized"
    audits = wiring["client"]._store["audit_logs"].values()
    assert sum(event["action"] == "generated_report.exported" for event in audits) == 3
